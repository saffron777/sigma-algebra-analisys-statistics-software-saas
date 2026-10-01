import io

from openpyxl import load_workbook

from sigma.infrastructure.exporters.xlsx_exporter import XlsxExporter
from tests.unit.infrastructure.exporter_fixtures import GREEK, sample_document


def open_book(document):
    return load_workbook(io.BytesIO(XlsxExporter().export(document)))


def test_sheets_and_order():
    book = open_book(sample_document())
    assert book.sheetnames[:2] == ["Índice", "Informe"]
    assert "T01 Coeficientes" in book.sheetnames and "T02 Tabla ANOVA" in book.sheetnames
    assert "T03 Tabla ANOVA" in book.sheetnames  # equal titles get distinct numbered sheets
    assert book.sheetnames[-2:] == ["Gráficos", "Datos ventas"]
    assert all(len(name) <= 31 for name in book.sheetnames)


def test_table_values_are_numbers_and_text_is_never_a_formula():
    book = open_book(sample_document())
    sheet = book["T01 Coeficientes"]
    header = [c.value for c in sheet[4]]
    assert header[:3] == ["Término", "Estimación", "Error est."]
    row = {c.column: c for c in sheet[5]}
    assert row[2].value == 12.3456 and isinstance(row[2].value, float)  # "12.3456" became a number
    assert row[5].value == "<0.001" and row[6].value == "[9.9, 14.7]"  # formatted text stays text
    hostile = {c.column: c for c in sheet[7]}
    assert hostile[1].value == "=1+1" and hostile[1].data_type == "s"  # a literal, not a formula
    assert hostile[2].value == "@cmd" and hostile[3].value == "+SUM(A1)"
    assert hostile[4].value == "—"  # None shows as a dash
    assert sheet.freeze_panes == "A5"


def test_narrative_sheet_has_text_tiles_alert_and_links():
    book = open_book(sample_document())
    sheet = book["Informe"]
    values = [c.value for r in sheet.iter_rows() for c in r if c.value is not None]
    assert "Regresión lineal múltiple: ventas ~ precio + publicidad" in values
    assert "R²" in values and "0.8421" in values and f"Símbolos: {GREEK}" in values
    assert "Hay 3 observaciones influyentes." in values
    assert any(isinstance(v, str) and v.startswith("Tabla: Coeficientes") for v in values)
    assert any(isinstance(v, str) and v.startswith("Gráfico 1") for v in values)
    links = [c.hyperlink for r in sheet.iter_rows() for c in r if c.hyperlink]
    assert links and all(link.location for link in links)


def test_charts_sheet_embeds_the_images():
    book = open_book(sample_document())
    assert len(book["Gráficos"]._images) == 2


def test_data_sheet_has_every_row_with_native_types():
    book = open_book(sample_document())
    sheet = book["Datos ventas"]
    assert sheet.max_row == 51 and sheet["A1"].value == "id" and sheet["B3"].value == 3.0
    assert sheet["D2"].value == "=cmd" and sheet["D2"].data_type == "s"
    assert sheet["C2"].value is None  # missing → empty cell
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:D51"


def test_index_carries_the_cover_data_and_branding():
    book = open_book(sample_document(branded=True))
    sheet = book["Índice"]
    texts = [c.value for r in sheet.iter_rows() for c in r if c.value]
    assert "Informe de ventas" in texts and "Ana Pérez" in texts and "Acme Estadística" in texts
    assert "2026-10-01" in texts and "Confidencial · Acme S.A." in texts
    assert len(sheet._images) == 1  # the logo
    assert sheet["A1"].font.color.rgb.endswith("DC2626")  # the title uses the brand colour…
    header = book["T01 Coeficientes"]["A4"]
    assert header.fill.start_color.rgb.endswith("DC2626")  # …and table headers the primary colour


def test_missing_images_are_reported_not_fatal():
    book = open_book(sample_document(with_images=False, data=False))
    assert "Gráficos" not in book.sheetnames
    values = [c.value for r in book["Informe"].iter_rows() for c in r if c.value]
    assert any("no disponible" in str(v) for v in values)
