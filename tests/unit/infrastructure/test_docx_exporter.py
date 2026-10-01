import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from sigma.domain.documents import LayoutOptions
from sigma.infrastructure.exporters.docx_exporter import DocxExporter
from tests.unit.infrastructure.exporter_fixtures import GREEK, sample_document


def open_doc(document):
    return Document(io.BytesIO(DocxExporter().export(document)))


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [c.text for row in table.rows for c in row.cells]
    return "\n".join(parts)


def test_structure_headings_and_content():
    doc = open_doc(sample_document())
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Contenido", "Regresión lineal múltiple: ventas ~ precio + publicidad", "ANOVA de un factor"]
    text = all_text(doc)
    assert "Informe de ventas" in text and "Ana Pérez" in text and "2026-10-01" in text  # cover
    assert "Resumen del modelo" in text and "0.8421" in text and f"Símbolos: {GREEK}" in text
    assert "Hay 3 observaciones influyentes." in text and "Revisar outliers" in text
    assert "df = sigma.load('ventas')" in text and "count  300.0" in text
    assert "Errores estándar clásicos." in text


def test_tables_have_repeating_shaded_headers_and_right_aligned_numbers():
    doc = open_doc(sample_document(branded=True))
    coefficients = next(t for t in doc.tables if t.rows[0].cells[0].text == "Término")
    header = coefficients.rows[0]
    assert header._tr.trPr.find(qn("w:tblHeader")) is not None  # repeats on every page
    shading = header.cells[0]._tc.tcPr.find(qn("w:shd"))
    assert shading.get(qn("w:fill")) == "DC2626"  # the template's primary colour
    assert coefficients.rows[1].cells[1].paragraphs[0].alignment == 2  # RIGHT
    assert [c.text for c in coefficients.rows[3].cells][:4] == ["=1+1", "@cmd", "+SUM(A1)", "—"]
    good = coefficients.rows[1].cells[4].paragraphs[0].runs[0]
    assert str(good.font.color.rgb) == "15803D"


def test_images_header_footer_and_branding():
    doc = open_doc(sample_document(branded=True))
    assert len(doc.inline_shapes) == 3  # two charts + the logo on the cover (the header holds another)
    section = doc.sections[0]
    header = " ".join(p.text for p in section.header.paragraphs)
    assert "Acme Estadística" in header and "Informe de ventas" in header
    footer_xml = section.footer._element.xml
    assert "Confidencial · Acme S.A." in footer_xml and "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert section.different_first_page_header_footer  # the cover has no header


def test_layout_options():
    doc = open_doc(
        sample_document(layout=LayoutOptions(page_size="Letter", orientation="landscape", cover=False, toc=False))
    )
    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE and section.page_width > section.page_height
    assert round(section.page_width.mm) == 279
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert "Contenido" not in headings and "Informe de ventas" not in all_text(doc).split("\n")[0]


def test_toc_field_and_update_on_open():
    doc = open_doc(sample_document())
    xml = doc.element.xml
    assert "TOC" in xml and "fldChar" in xml
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_missing_images_do_not_break_the_export():
    doc = open_doc(sample_document(with_images=False))
    assert "(Gráfico no disponible)" in all_text(doc)


def test_core_properties():
    doc = open_doc(sample_document())
    assert doc.core_properties.title == "Informe de ventas" and doc.core_properties.author == "Ana Pérez"
