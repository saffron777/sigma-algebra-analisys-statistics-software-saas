"""Exporting (PDF / Word / Excel) and report templates, through the real HTTP endpoints."""

import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.integration.conftest import import_frame
from tests.unit.infrastructure.exporter_fixtures import make_logo, make_png

FORMATS = ["pdf", "docx", "xlsx"]
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def result_sections():
    return [
        {
            "title": "Regresión de ventas",
            "subtitle": "ventas ~ precio",
            "meta": ["Dataset: ventas"],
            "blocks": [
                {"type": "heading", "text": "Modelo"},
                {"type": "tiles", "items": [{"label": "R²", "value": "0.84"}]},
                {
                    "type": "table",
                    "title": "Coeficientes",
                    "columns": ["Término", "β"],
                    "rows": [["precio", "-2.5"], ["=1+1", {"text": "0.4", "tone": "bad"}]],
                },
                {"type": "figure", "index": 0, "caption": "Ajuste Σ"},
            ],
        },
        {"title": "ANOVA", "blocks": [{"type": "text", "paragraphs": ["Hay diferencias."]}]},
    ]


def export(client, fmt, *, payload=None, images=None, template=None, **extra):
    body = {"title": "Informe de prueba", "author": "Ana", "sections": result_sections(), **extra}
    if template:
        body["template_id"] = template
    body.update(payload or {})
    files = {f"img-{s}-{i}": (io.BytesIO(png), f"g{s}{i}.png") for (s, i), png in (images or {}).items() if png}
    return client.post(
        f"/api/export/{fmt}",
        data={"payload": json.dumps(body), **files},
        content_type="multipart/form-data",
    )


def create_template(client, **fields):
    response = client.post("/api/templates", {"name": "Corporativa", **fields})
    assert response.status_code == 201, response.get_data(as_text=True)
    return response.parsed["template"]


# ------------------------------------------------------------------ exporting a report
@pytest.mark.parametrize("fmt", FORMATS)
def test_every_format_downloads_a_valid_file(api, fmt):
    response = export(api, fmt, images={(0, 0): make_png()})
    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.mimetype == CONTENT_TYPES[fmt]
    disposition = response.headers["Content-Disposition"]
    assert disposition.startswith("attachment;") and f".{fmt}" in disposition and "filename*=UTF-8''" in disposition
    assert "Sigma - Informe de prueba - " in disposition
    assert response.headers["X-Content-Type-Options"] == "nosniff" and "no-store" in response.headers["Cache-Control"]
    data = response.get_data()
    if fmt == "pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(p.extract_text() for p in reader.pages)
        assert "Informe de prueba" in text and "Regresión de ventas" in text and "Hay diferencias." in text
        assert "Ana" in text and sum(len(p.images) for p in reader.pages) >= 1
    elif fmt == "docx":
        doc = Document(io.BytesIO(data))
        assert "Regresión de ventas" in [p.text for p in doc.paragraphs] and len(doc.inline_shapes) == 1
    else:
        book = load_workbook(io.BytesIO(data))
        assert "T01 Coeficientes" in book.sheetnames and len(book["Gráficos"]._images) == 1
        assert book["T01 Coeficientes"]["A6"].value == "=1+1"


def test_json_body_without_images_also_works(api):
    response = api.post("/api/export/pdf", {"title": "Solo texto", "sections": result_sections()})
    assert response.status_code == 200 and response.get_data().startswith(b"%PDF")


def test_a_single_section_uses_its_title_for_the_file(api):
    response = api.post("/api/export/xlsx", {"sections": result_sections()[:1]})
    assert "Sigma - Regresión de ventas - " in response.headers["Content-Disposition"].split("filename*=UTF-8''")[
        1
    ].replace("%20", " ").replace("%C3%B3", "ó")


def test_export_validation(api):
    assert export(api, "odt").status_code == 422
    assert api.post("/api/export/pdf", {"sections": []}).status_code == 422
    assert api.post("/api/export/pdf", {"sections": [{"title": "  "}]}).status_code == 422
    assert api.post("/api/export/pdf", {"sections": result_sections(), "page_size": "A0"}).status_code == 422
    assert api.post("/api/export/pdf", {"sections": result_sections(), "orientation": "diagonal"}).status_code == 422
    assert api.post("/api/export/pdf", {"sections": result_sections(), "template_id": "0" * 32}).status_code == 404
    bad_json = api.post("/api/export/pdf", data={"payload": "{no"}, content_type="multipart/form-data")
    assert bad_json.status_code == 422


def test_chart_images_are_validated(api):
    not_png = export(api, "pdf", images={(0, 0): b"GIF89a....."})
    assert not_png.status_code == 422 and "PNG" in not_png.parsed["error"]["message"]
    broken = export(api, "pdf", images={(0, 0): b"\x89PNG\r\n\x1a\n" + b"garbage" * 10})
    assert broken.status_code == 422
    too_big = export(api, "pdf", images={(0, 0): b"\x89PNG\r\n\x1a\n" + b"0" * (9 * 1024 * 1024)})
    assert too_big.status_code == 422 and "8 MB" in too_big.parsed["error"]["message"]


def test_hostile_content_is_rendered_as_text_not_executed(api):
    sections = [
        {
            "title": "<script>x</script>",
            "blocks": [{"type": "text", "paragraphs": ["<b>hola</b> &amp;"]}, {"type": "evil", "x": 1}],
        }
    ]
    response = api.post("/api/export/docx", {"sections": sections})
    assert response.status_code == 200
    doc = Document(io.BytesIO(response.get_data()))
    assert "<b>hola</b> &amp;" in [p.text for p in doc.paragraphs]


def test_options_override_the_template(api):
    template = create_template(api, show_cover=True, show_toc=True)
    plain = export(
        api,
        "pdf",
        template=template["id"],
        payload={"cover": False, "toc": False, "orientation": "landscape", "page_size": "Letter"},
    )
    reader = PdfReader(io.BytesIO(plain.get_data()))
    assert "Contenido" not in reader.pages[0].extract_text() and round(float(reader.pages[0].mediabox.width)) == 792


# ------------------------------------------------------------------ templates
def test_template_crud_and_default(api):
    first = create_template(api, organization="Acme", primary_color="#DC2626", footer_text="Confidencial")
    assert first["is_default"] is True and first["primary_color"] == "#dc2626" and first["has_logo"] is False
    second = create_template(api, name="Interna", page_size="Letter", orientation="landscape", show_cover=False)
    assert second["is_default"] is False and second["show_cover"] is False
    assert [t["name"] for t in api.get("/api/templates").parsed["templates"]] == [
        "Corporativa",
        "Interna",
    ]  # default first
    assert (
        api.post(f"/api/templates/{second['id']}/default", {"default": True}).parsed["template"]["is_default"] is True
    )
    listed = {t["name"]: t["is_default"] for t in api.get("/api/templates").parsed["templates"]}
    assert listed == {"Corporativa": False, "Interna": True}  # only one default at a time
    edited = api.put(
        f"/api/templates/{second['id']}", {"name": "Interna 2", "organization": "Acme SA", "primary_color": "#112233"}
    )
    assert edited.parsed["template"]["name"] == "Interna 2" and edited.parsed["template"]["primary_color"] == "#112233"
    assert api.delete(f"/api/templates/{second['id']}").status_code == 200
    assert [t["is_default"] for t in api.get("/api/templates").parsed["templates"]] == [
        True
    ]  # the remaining one inherits it
    assert api.get(f"/api/templates/{second['id']}").status_code == 404


def test_template_validation(api):
    create_template(api)
    assert (
        api.post("/api/templates", {"name": "corporativa"}).status_code == 422
    )  # names are unique, case-insensitively
    for bad in (
        {"name": ""},
        {"name": "x", "primary_color": "rojo"},
        {"name": "x", "page_size": "A0"},
        {"name": "x", "orientation": "z"},
        {"name": "x" * 90},
    ):
        assert api.post("/api/templates", bad).status_code == 422, bad


def test_logo_upload_is_validated_normalised_and_served(api):
    template = create_template(api)
    tid = template["id"]

    def upload(content, name="logo.png"):
        return api.post(
            f"/api/templates/{tid}/logo", data={"file": (io.BytesIO(content), name)}, content_type="multipart/form-data"
        )

    ok = upload(make_logo())
    assert ok.status_code == 201 and ok.parsed["template"]["has_logo"] is True
    served = api.get(f"/api/templates/{tid}/logo")
    assert (
        served.status_code == 200
        and served.mimetype == "image/png"
        and served.headers["X-Content-Type-Options"] == "nosniff"
    )
    assert served.get_data().startswith(b"\x89PNG")
    svg = upload(b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', "x.svg")
    assert svg.status_code == 422 and "PNG o JPG" in svg.parsed["error"]["message"]
    assert upload(b"").status_code == 422
    assert upload(b"\x89PNG\r\n\x1a\n" + b"x" * 3_000_000).status_code == 422  # over 2 MB
    assert upload(b"not an image at all").status_code == 422
    assert api.post(f"/api/templates/{tid}/logo", data={}, content_type="multipart/form-data").status_code == 422
    assert api.delete(f"/api/templates/{tid}/logo").parsed["template"]["has_logo"] is False
    assert api.get(f"/api/templates/{tid}/logo").status_code == 404


def test_large_logos_are_downsized_and_jpeg_is_accepted(api):
    from PIL import Image

    tid = create_template(api)["id"]
    buffer = io.BytesIO()
    Image.new("RGB", (3000, 1200), (10, 120, 200)).save(buffer, format="JPEG")
    assert (
        api.post(
            f"/api/templates/{tid}/logo",
            data={"file": (io.BytesIO(buffer.getvalue()), "a.jpg")},
            content_type="multipart/form-data",
        ).status_code
        == 201
    )
    with Image.open(io.BytesIO(api.get(f"/api/templates/{tid}/logo").get_data())) as image:
        assert image.format == "PNG" and image.width <= 600 and image.height <= 240


def test_template_branding_reaches_the_documents(api):
    template = create_template(
        api,
        organization="Acme Estadística",
        primary_color="#dc2626",
        footer_text="Confidencial · Acme",
        cover_text="Para la gerencia",
    )
    api.post(
        f"/api/templates/{template['id']}/logo",
        data={"file": (io.BytesIO(make_logo()), "l.png")},
        content_type="multipart/form-data",
    )
    pdf = PdfReader(io.BytesIO(export(api, "pdf", template=template["id"]).get_data()))
    cover = pdf.pages[0].extract_text()
    assert "Acme Estadística" in cover and "Para la gerencia" in cover and pdf.pages[0].images
    assert "Confidencial · Acme" in pdf.pages[2].extract_text()
    docx = Document(io.BytesIO(export(api, "docx", template=template["id"]).get_data()))
    assert "Acme Estadística" in " ".join(p.text for p in docx.sections[0].header.paragraphs)
    book = load_workbook(io.BytesIO(export(api, "xlsx", template=template["id"]).get_data()))
    assert book["T01 Coeficientes"]["A4"].fill.start_color.rgb.endswith("DC2626")
    # the default template applies without naming it
    by_default = PdfReader(io.BytesIO(export(api, "pdf").get_data()))
    assert "Acme Estadística" in by_default.pages[0].extract_text()


def test_templates_are_private_per_account(api, make_client):
    mine = create_template(api)
    api.post(
        f"/api/templates/{mine['id']}/logo",
        data={"file": (io.BytesIO(make_logo()), "l.png")},
        content_type="multipart/form-data",
    )
    other = make_client("luis")
    assert other.get("/api/templates").parsed["templates"] == []
    for response in (
        other.get(f"/api/templates/{mine['id']}"),
        other.put(f"/api/templates/{mine['id']}", {"name": "robada"}),
        other.delete(f"/api/templates/{mine['id']}"),
        other.get(f"/api/templates/{mine['id']}/logo"),
        other.post(f"/api/templates/{mine['id']}/default", {"default": True}),
        other.delete(f"/api/templates/{mine['id']}/logo"),
        export(other, "pdf", template=mine["id"]),
    ):
        assert response.status_code == 404, response.request.path
    assert other.post("/api/templates", {"name": "Corporativa"}).status_code == 201  # same name, other owner
    assert api.get(f"/api/templates/{mine['id']}").parsed["template"]["has_logo"] is True


def test_templates_show_up_in_sync(api):
    before = api.get("/api/sync").parsed["templates"]
    create_template(api)
    assert api.get("/api/sync").parsed["templates"] > before


# ------------------------------------------------------------------ dataset ficha
@pytest.mark.parametrize("fmt", FORMATS)
def test_dataset_ficha_download(api, sales_frame, fmt):
    dataset = import_frame(api, sales_frame, "ventas")
    response = api.get(f"/api/datasets/{dataset['id']}/report/{fmt}?author=Ana&cover=0")
    assert response.status_code == 200 and response.mimetype == CONTENT_TYPES[fmt]
    assert "Dataset ventas" in response.headers["Content-Disposition"]
    data = response.get_data()
    if fmt == "pdf":
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
        assert "Ficha del dataset" in text and "precio" in text and "Primeras 20 filas" in text
    elif fmt == "docx":
        assert any("Ficha del dataset" in p.text for p in Document(io.BytesIO(data)).paragraphs)
    else:
        book = load_workbook(io.BytesIO(data))
        assert book["Datos"].max_row == 201  # EVERY row, not just the preview
        assert book["Datos"]["A1"].value == "ventas" and isinstance(book["Datos"]["B2"].value, float)


def test_dataset_ficha_validation_and_ownership(api, make_client, sales_frame):
    dataset = import_frame(api, sales_frame, "ventas")
    assert api.get(f"/api/datasets/{dataset['id']}/report/odt").status_code == 422
    assert api.get(f"/api/datasets/{'0' * 32}/report/pdf").status_code == 404
    other = make_client("luis")
    assert other.get(f"/api/datasets/{dataset['id']}/report/pdf").status_code == 404


# ------------------------------------------------------------------ notebook report
@pytest.mark.parametrize("fmt", FORMATS)
def test_notebook_report_download(api, fmt):
    created = api.post("/api/notebooks", {"name": "Mi análisis"}).parsed["notebook"]
    cells = [
        {"id": "a1", "cell_type": "markdown", "source": "# Hallazgos\nTexto **clave**.", "metadata": {}},
        {
            "id": "a2",
            "cell_type": "code",
            "source": "print(40 + 2)",
            "execution_count": 1,
            "metadata": {"sigma": {"page_break_before": True}},
            "outputs": [{"output_type": "stream", "name": "stdout", "text": "42\n"}],
        },
    ]
    assert api.put(f"/api/notebooks/{created['id']}", {"cells": cells}).status_code == 200
    response = api.get(f"/api/notebooks/{created['id']}/report/{fmt}?numbers=1")
    assert response.status_code == 200 and response.mimetype == CONTENT_TYPES[fmt]
    data = response.get_data()
    if fmt == "pdf":
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
        assert "Hallazgos" in text and "print(40 + 2)" in text and "42" in text
    elif fmt == "docx":
        text = "\n".join(c.text for t in Document(io.BytesIO(data)).tables for r in t.rows for c in r.cells)
        assert "print(40 + 2)" in text and "42" in text
    else:
        book = load_workbook(io.BytesIO(data))
        values = [c.value for r in book["Informe"].iter_rows() for c in r if c.value]
        assert "print(40 + 2)" in values and "42" in values
    only_text = api.get(f"/api/notebooks/{created['id']}/report/docx?code=0&outputs=0")
    assert "print(40 + 2)" not in "\n".join(
        c.text for t in Document(io.BytesIO(only_text.get_data())).tables for r in t.rows for c in r.cells
    )


def test_notebook_report_needs_the_console_permission_and_ownership(api, make_client):
    created = api.post("/api/notebooks", {"name": "x"}).parsed["notebook"]
    no_console = make_client("sin_consola")
    assert no_console.get(f"/api/notebooks/{created['id']}/report/pdf").status_code == 403
    other = make_client("con_consola", console=True)
    assert other.get(f"/api/notebooks/{created['id']}/report/pdf").status_code == 404
    assert api.get(f"/api/notebooks/{created['id']}/report/odt").status_code == 422
