"""Excel export (openpyxl).

Sheets: «Índice» (cover data + hyperlinks), «Informe» (the narrative: text, key figures, alerts, code and links to the
tables and charts), one sheet per table (numbers stored as numbers, frozen header), «Gráficos» (every chart as an
image) and one sheet per ``DataTable`` (e.g. every row of a dataset).
"""

from __future__ import annotations

import contextlib
import datetime as dt
import io
import math
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from sigma.domain.documents import (
    DataTable,
    ReportDocument,
    cell_number,
    cell_text,
    cell_tone,
    normalize_color,
)
from sigma.domain.markdown_blocks import plain
from sigma.infrastructure.exporters.common import figure_image, flatten_png, image_size, tint

MAX_ROWS = 1_048_575
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_TONE_FONT = {"strong": "000000", "good": "15803D", "bad": "B91C1C", "muted": "64748B"}
_ALERT_FILL = {"info": "E0F2FE", "success": "DCFCE7", "warning": "FEF3C7", "danger": "FEE2E2"}
_THIN = Side(style="thin", color="CBD5E1")


def _argb(color: str) -> str:
    return normalize_color(color, "#000000").lstrip("#").upper()


def _set_text(cell, value: Any) -> None:
    """Write text so that it can never be interpreted as a formula (``=1+1`` stays literal text)."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _python(value: Any) -> Any:
    """Plain Python scalar for a cell (numpy numbers, NaN and timestamps are what dataset rows contain)."""
    if value is None:
        return None
    if hasattr(value, "item") and not isinstance(value, (dt.date, dt.datetime)):
        with contextlib.suppress(ValueError, AttributeError):
            value = value.item()
    if hasattr(value, "to_pydatetime"):
        value = value.to_pydatetime()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        value = value.astimezone(dt.UTC).replace(tzinfo=None)
    return value


class XlsxExporter:
    format = "xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def export(self, document: ReportDocument) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        branding = document.branding
        self._primary = _argb(branding.primary_color)
        self._header_fill = PatternFill("solid", start_color=self._primary)
        self._soft_fill = PatternFill("solid", start_color=_argb(tint(branding.primary_color, 0.88)))
        self._names: set[str] = set()

        index = self._sheet(workbook, "Índice")
        report = self._sheet(workbook, "Informe")
        table_sheets: list[tuple[str, str, str]] = []  # (sheet name, title, section)
        self._build_report(workbook, report, document, table_sheets)
        charts = self._build_charts(workbook, document)
        data_sheets = [self._build_data_sheet(workbook, table) for table in document.data_tables]
        self._build_index(index, document, table_sheets, charts, data_sheets)

        workbook.properties.title = document.title[:250]
        workbook.properties.creator = document.author or branding.organization or "Sigma"
        workbook.active = 0
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------ naming
    def _sheet(self, workbook: Workbook, name: str) -> Worksheet:
        cleaned = _SHEET_FORBIDDEN.sub(" ", name).strip().strip("'") or "Hoja"
        cleaned = cleaned[:31]
        candidate, counter = cleaned, 2
        while candidate.lower() in self._names:
            suffix = f" ({counter})"
            candidate = cleaned[: 31 - len(suffix)] + suffix
            counter += 1
        self._names.add(candidate.lower())
        return workbook.create_sheet(candidate)

    @staticmethod
    def _link(cell, sheet: str, text: str) -> None:
        cell.value = text
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1", display=text)
        cell.font = Font(color="0563C1", underline="single")

    # ------------------------------------------------------------------ narrative
    def _build_report(self, workbook, sheet: Worksheet, document: ReportDocument, table_sheets: list) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 30
        for letter in "BCDEFGH":
            sheet.column_dimensions[letter].width = 18
        row = 1
        figure_count = 0
        for section in document.sections:
            cell = sheet.cell(row=row, column=1, value=section.title)
            cell.font = Font(size=15, bold=True, color="FFFFFF")
            for column in range(1, 9):
                sheet.cell(row=row, column=column).fill = self._header_fill
            row += 1
            if section.subtitle:
                _set_text(sheet.cell(row=row, column=1), section.subtitle)
                sheet.cell(row=row, column=1).font = Font(italic=True, color="475569")
                row += 1
            for line in section.meta:
                _set_text(sheet.cell(row=row, column=1), line)
                sheet.cell(row=row, column=1).font = Font(size=9, color="64748B")
                row += 1
            row += 1
            for block in section.blocks:
                kind = block.get("type")
                if kind == "heading":
                    cell = sheet.cell(row=row, column=1)
                    _set_text(cell, block.get("text", ""))
                    cell.font = Font(size=12, bold=True, color=self._primary)
                    row += 1
                elif kind == "tiles":
                    for item in block.get("items", []):
                        label = sheet.cell(row=row, column=1)
                        _set_text(label, item.get("label", ""))
                        label.font = Font(bold=True)
                        label.fill = self._soft_fill
                        _set_text(sheet.cell(row=row, column=2), item.get("value", ""))
                        if item.get("hint"):
                            _set_text(sheet.cell(row=row, column=3), item["hint"])
                            sheet.cell(row=row, column=3).font = Font(color="64748B", size=9)
                        row += 1
                    row += 1
                elif kind in ("text", "list"):
                    paragraphs = (
                        [plain(p) if block.get("markdown") else p for p in block.get("paragraphs", [])]
                        if kind == "text"
                        else [
                            f"{'•' if not block.get('ordered') else str(i + 1) + '.'} {plain(t)}"
                            for i, t in enumerate(block.get("items", []))
                        ]
                    )
                    for paragraph in paragraphs:
                        row = self._wrapped(sheet, row, paragraph)
                    row += 1
                elif kind == "alert":
                    cell = sheet.cell(row=row, column=1)
                    _set_text(cell, str(block.get("message", "")))
                    self._merge_wrap(sheet, row, str(block.get("message", "")))
                    for column in range(1, 9):
                        sheet.cell(row=row, column=column).fill = PatternFill(
                            "solid", start_color=_ALERT_FILL.get(block.get("kind"), "E0F2FE")
                        )
                    row += 2
                elif kind in ("code", "output"):
                    text = block.get("source") if kind == "code" else block.get("text", "")
                    cell = sheet.cell(row=row, column=1)
                    _set_text(cell, str(text or ""))
                    cell.font = Font(name="Consolas", size=9)
                    self._merge_wrap(sheet, row, str(text or ""), chars=110)
                    for column in range(1, 9):
                        sheet.cell(row=row, column=column).fill = PatternFill("solid", start_color="F1F5F9")
                    row += 2
                elif kind == "table":
                    name = self._table_sheet(workbook, section, block, len(table_sheets) + 1)
                    table_sheets.append((name, block.get("title") or section.title, section.title))
                    cell = sheet.cell(row=row, column=1)
                    self._link(cell, name, f"Tabla: {block.get('title') or 'ver hoja'} → hoja «{name}»")
                    row += 2
                elif kind == "figure":
                    figure_count += 1
                    cell = sheet.cell(row=row, column=1)
                    if figure_image(section, block):
                        cell.value = (
                            f"Gráfico {figure_count}: {block.get('caption') or section.title} (ver hoja «Gráficos»)"
                        )
                    else:
                        cell.value = f"Gráfico {figure_count}: {block.get('caption') or ''} (no disponible)"
                    cell.font = Font(italic=True, color="64748B")
                    row += 2
            row += 2

    @staticmethod
    def _merge_wrap(sheet: Worksheet, row: int, text: str, columns: int = 8, chars: int = 120) -> None:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
        cell = sheet.cell(row=row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        lines = sum(max(1, math.ceil(len(part) / chars)) for part in text.split("\n"))
        sheet.row_dimensions[row].height = min(409, max(15, 15 * lines))

    def _wrapped(self, sheet: Worksheet, row: int, text: str) -> int:
        _set_text(sheet.cell(row=row, column=1), text)
        self._merge_wrap(sheet, row, text)
        return row + 1

    # ------------------------------------------------------------------ one sheet per table
    def _table_sheet(self, workbook, section, block: dict[str, Any], number: int) -> str:
        title = block.get("title") or section.title
        sheet = self._sheet(workbook, f"T{number:02d} {title}")
        title_cell = sheet.cell(row=1, column=1)
        _set_text(title_cell, str(title))
        title_cell.font = Font(size=13, bold=True, color=self._primary)
        if block.get("note"):
            _set_text(sheet.cell(row=2, column=1), str(block["note"]))
            sheet.cell(row=2, column=1).font = Font(size=9, italic=True, color="64748B")
        columns = [str(c) for c in block.get("columns", [])]
        self._write_table(sheet, 4, columns, block.get("rows", []))
        sheet.freeze_panes = sheet.cell(row=5, column=1)
        back = sheet.cell(row=3, column=1)
        self._link(back, "Índice", "← Índice")
        return sheet.title

    def _write_table(self, sheet: Worksheet, first_row: int, columns: list[str], rows: list[list[Any]]) -> None:
        widths = [len(c) for c in columns]
        for offset, name in enumerate(columns, start=1):
            cell = sheet.cell(row=first_row, column=offset)
            _set_text(cell, name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = self._header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows[: MAX_ROWS - first_row], start=first_row + 1):
            for c, value in enumerate(row[: len(columns)], start=1):
                cell = sheet.cell(row=r, column=c)
                number = cell_number(value)
                text = cell_text(value)
                if number is not None:
                    cell.value = number
                    cell.alignment = Alignment(horizontal="right")
                else:
                    _set_text(cell, text)
                tone = cell_tone(value)
                if tone:
                    cell.font = Font(bold=tone == "strong", color=_TONE_FONT[tone])
                cell.border = Border(bottom=_THIN)
                widths[c - 1] = max(widths[c - 1], len(text))
        for offset, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(offset)].width = min(60, max(10, width + 2))

    # ------------------------------------------------------------------ whole data tables
    def _build_data_sheet(self, workbook, table: DataTable) -> str:
        sheet = self._sheet(workbook, table.name)
        widths = [len(str(c)) for c in table.columns]
        for offset, name in enumerate(table.columns, start=1):
            cell = sheet.cell(row=1, column=offset)
            _set_text(cell, str(name))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = self._header_fill
        for r, row in enumerate(table.rows[: MAX_ROWS - 1], start=2):
            for c, raw in enumerate(row, start=1):
                value = _python(raw)
                cell = sheet.cell(row=r, column=c)
                if isinstance(value, dt.datetime):
                    cell.value = value
                    cell.number_format = "yyyy-mm-dd hh:mm:ss" if value.time() != dt.time(0) else "yyyy-mm-dd"
                elif isinstance(value, str):
                    _set_text(cell, value)
                else:
                    cell.value = value
                if r <= 200:  # widths are estimated from the first rows only (keeps big tables fast)
                    widths[c - 1] = max(widths[c - 1], len(str(value)) if value is not None else 0)
        for offset, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(offset)].width = min(50, max(10, width + 2))
        sheet.freeze_panes = "A2"
        if table.rows:
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(max(len(table.columns), 1))}{min(len(table.rows) + 1, MAX_ROWS)}"
            )
        return sheet.title

    # ------------------------------------------------------------------ charts
    def _build_charts(self, workbook, document: ReportDocument) -> str | None:
        placed = [
            (section, block, figure_image(section, block))
            for section in document.sections
            for block in section.blocks
            if block.get("type") == "figure"
        ]
        placed = [item for item in placed if item[2]]
        if not placed:
            return None
        sheet = self._sheet(workbook, "Gráficos")
        sheet.sheet_view.showGridLines = False
        row = 1
        for number, (section, block, png) in enumerate(placed, start=1):
            title = sheet.cell(row=row, column=1)
            _set_text(title, f"Gráfico {number} — {block.get('caption') or section.title}")
            title.font = Font(bold=True, color=self._primary)
            size = image_size(png)
            if size is None:
                row += 2
                continue
            picture = XlImage(io.BytesIO(flatten_png(png)))
            scale = min(1.0, 820 / size[0])
            picture.width, picture.height = round(size[0] * scale), round(size[1] * scale)
            sheet.add_image(picture, f"A{row + 1}")
            row += math.ceil(picture.height / 20) + 3
        return sheet.title

    # ------------------------------------------------------------------ index
    def _build_index(self, sheet: Worksheet, document: ReportDocument, tables: list, charts, data_sheets) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 70
        title = sheet.cell(row=1, column=1)
        _set_text(title, document.title)
        title.font = Font(size=18, bold=True, color=self._primary)
        row = 2
        for label, value in (
            ("Subtítulo", document.subtitle),
            ("Organización", document.branding.organization),
            ("Autor", document.author),
            ("Fecha", document.created_at[:10] if document.created_at else ""),
        ):
            if value:
                _set_text(sheet.cell(row=row, column=1), label)
                sheet.cell(row=row, column=1).font = Font(bold=True)
                _set_text(sheet.cell(row=row, column=2), value)
                row += 1
        if document.branding.logo_png:
            size = image_size(document.branding.logo_png)
            if size:
                logo = XlImage(io.BytesIO(flatten_png(document.branding.logo_png)))
                scale = min(1.0, 90 / size[1])
                logo.width, logo.height = round(size[0] * scale), round(size[1] * scale)
                sheet.add_image(logo, "D1")
        row += 1
        sheet.cell(row=row, column=1, value="Contenido").font = Font(bold=True, size=12, color=self._primary)
        row += 1
        self._link(sheet.cell(row=row, column=1), "Informe", "Informe")
        _set_text(sheet.cell(row=row, column=2), "Resumen narrativo: texto, cifras clave, advertencias")
        row += 1
        for name, table_title, section_title in tables:
            self._link(sheet.cell(row=row, column=1), name, name)
            _set_text(sheet.cell(row=row, column=2), f"{table_title} — {section_title}"[:200])
            row += 1
        if charts:
            self._link(sheet.cell(row=row, column=1), charts, charts)
            _set_text(sheet.cell(row=row, column=2), "Todos los gráficos como imágenes")
            row += 1
        for name in data_sheets:
            self._link(sheet.cell(row=row, column=1), name, name)
            _set_text(sheet.cell(row=row, column=2), "Datos completos")
            row += 1
        if document.branding.footer_text:
            row += 1
            _set_text(sheet.cell(row=row, column=1), document.branding.footer_text)
            sheet.cell(row=row, column=1).font = Font(size=9, color="64748B")
