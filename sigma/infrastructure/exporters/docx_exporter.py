"""Word export (python-docx): cover, table of contents field, one section per result, branded header and footer."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor

from sigma.domain.documents import (
    ReportDocument,
    ReportSection,
    cell_text,
    cell_tone,
    is_numeric_text,
    normalize_color,
)
from sigma.domain.markdown_blocks import plain
from sigma.infrastructure.exporters.common import (
    document_date,
    figure_image,
    flatten_png,
    hex_to_rgb,
    image_size,
    paragraph_runs,
    tint,
)

_PAGES_MM = {"A4": (210, 297), "Letter": (215.9, 279.4), "Legal": (215.9, 355.6)}
_TONE_COLOR = {"strong": "000000", "good": "15803D", "bad": "B91C1C", "muted": "64748B"}
_ALERT = {  # (background, left bar)
    "info": ("E0F2FE", "0284C7"),
    "success": ("DCFCE7", "16A34A"),
    "warning": ("FEF3C7", "D97706"),
    "danger": ("FEE2E2", "DC2626"),
}
_FONT = "Calibri"
_MONO = "Consolas"


def _rgb(color: str) -> RGBColor:
    return RGBColor(*hex_to_rgb(color))


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill.lstrip("#").upper())
    properties.append(shading)


def _borders(cell, **edges: tuple[str, int]) -> None:
    """``left=("D97706", 24)``: a coloured edge of the given width (1/8 pt)."""
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:tcBorders")
    for edge, (color, size) in edges.items():
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)
        element.append(node)
    shading = properties.find(qn("w:shd"))
    if shading is not None:  # schema order: borders come before shading
        shading.addprevious(element)
    else:
        properties.append(element)


def _field(paragraph, instruction: str, placeholder: str = "1") -> None:
    """A Word field (PAGE, NUMPAGES, TOC…) that Word fills in when it opens or prints the document."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, code, separate])
    paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    properties.append(node)


def _no_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


class DocxExporter:
    format = "docx"
    content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extension = "docx"

    def export(self, document: ReportDocument) -> bytes:
        self._doc = Document()
        self._brand = document.branding
        self._primary = normalize_color(self._brand.primary_color, "#4f46e5")
        self._secondary = normalize_color(self._brand.secondary_color, "#0f172a")
        self._setup_page(document)
        self._setup_styles()
        self._header_footer(document)
        if document.layout.cover:
            self._cover(document)
        if document.layout.toc and document.sections:
            self._toc(document)
        for position, section in enumerate(document.sections):
            if position or document.layout.cover or document.layout.toc:
                self._doc.add_page_break()
            self._section(section)
        properties = self._doc.core_properties
        properties.title = document.title[:250]
        properties.author = document.author or self._brand.organization or "Sigma"
        properties.subject = document.subtitle or ""
        self._ask_to_update_fields()
        buffer = io.BytesIO()
        self._doc.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------ page and styles
    def _setup_page(self, document: ReportDocument) -> None:
        section = self._doc.sections[0]
        width, height = _PAGES_MM.get(document.layout.page_size, _PAGES_MM["A4"])
        landscape = document.layout.orientation == "landscape"
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (Mm(height), Mm(width)) if landscape else (Mm(width), Mm(height))
        margin = Mm(max(10, min(40, document.layout.margin_mm)))
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = margin
        section.header_distance = section.footer_distance = Mm(8)
        self._width = section.page_width - section.left_margin - section.right_margin  # EMU

    def _setup_styles(self) -> None:
        styles = self._doc.styles
        normal = styles["Normal"]
        normal.font.name = _FONT
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)
        normal.paragraph_format.space_after = Pt(6)
        for name, size, color in (
            ("Heading 1", 18, self._primary),
            ("Heading 2", 13.5, self._secondary),
            ("Heading 3", 11.5, self._secondary),
            ("Heading 4", 10.5, self._secondary),
        ):
            style = styles[name]
            style.font.name = _FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = _rgb(color)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)
            style.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True

    def _ask_to_update_fields(self) -> None:
        """Make Word refresh the table of contents (page numbers) the first time the file is opened."""
        settings = self._doc.settings.element
        node = OxmlElement("w:updateFields")
        node.set(qn("w:val"), "true")
        settings.append(node)

    # ------------------------------------------------------------------ header / footer / cover
    def _header_footer(self, document: ReportDocument) -> None:
        section = self._doc.sections[0]
        section.different_first_page_header_footer = bool(document.layout.cover)
        header = section.header.paragraphs[0]
        header.text = ""
        if self._brand.logo_png and (size := image_size(self._brand.logo_png)):
            height = Mm(9)
            header.add_run().add_picture(
                io.BytesIO(flatten_png(self._brand.logo_png)), width=Emu(int(height * size[0] / size[1])), height=height
            )
            header.add_run("  ")
        if self._brand.organization:
            run = header.add_run(self._brand.organization)
            run.bold = True
            run.font.color.rgb = _rgb(self._primary)
        right = header.add_run("\t" + document.title[:70])
        right.font.size = Pt(8.5)
        right.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        self._tab_stop(header)
        footer = section.footer.paragraphs[0]
        footer.text = ""
        note = footer.add_run((self._brand.footer_text or "Generado con Sigma") + "\t")
        note.font.size = Pt(8.5)
        note.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        for run_text, instruction in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
            label = footer.add_run(run_text)
            label.font.size = Pt(8.5)
            _field(footer, instruction)
        self._tab_stop(footer)

    def _tab_stop(self, paragraph) -> None:
        properties = paragraph._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        for kind, position in (("clear", 4513), ("clear", 9026), ("right", int(self._width / 635))):
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), kind)
            tab.set(qn("w:pos"), str(position))
            tabs.append(tab)
        properties.append(tabs)

    def _cover(self, document: ReportDocument) -> None:
        doc = self._doc
        if self._brand.logo_png and (size := image_size(self._brand.logo_png)):
            paragraph = doc.add_paragraph()
            height = Mm(24)
            paragraph.add_run().add_picture(
                io.BytesIO(flatten_png(self._brand.logo_png)), width=Emu(int(height * size[0] / size[1])), height=height
            )
        for _ in range(4):
            doc.add_paragraph()
        title = doc.add_paragraph()
        run = title.add_run(document.title)
        run.bold = True
        run.font.size = Pt(30)
        run.font.color.rgb = _rgb(self._primary)
        title.paragraph_format.space_after = Pt(8)
        if document.subtitle:
            sub = doc.add_paragraph().add_run(document.subtitle)
            sub.font.size = Pt(15)
            sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        rule = doc.add_paragraph()
        rule.paragraph_format.space_after = Pt(18)
        rule.add_run("▬▬▬").font.color.rgb = _rgb(self._primary)
        for line in (self._brand.organization, self._brand.cover_text):
            if line:
                doc.add_paragraph().add_run(line).font.size = Pt(12)
        meta = " · ".join(part for part in (document.author, document_date(document)) if part)
        if meta:
            run = doc.add_paragraph().add_run(meta)
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        if len(document.sections) > 1:
            note = doc.add_paragraph().add_run(f"{len(document.sections)} análisis")
            note.font.size = Pt(10)
            note.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        if document.layout.toc or document.sections:
            doc.add_page_break()

    def _toc(self, document: ReportDocument) -> None:
        self._doc.add_heading("Contenido", level=1)
        paragraph = self._doc.add_paragraph()
        _field(
            paragraph,
            'TOC \\o "1-2" \\h \\z \\u',
            "Haga clic derecho aquí y elija «Actualizar campo» para ver el índice.",
        )
        for position, section in enumerate(document.sections, start=1):
            entry = self._doc.add_paragraph(f"{position}. {section.title}")
            entry.paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------------ sections and blocks
    def _section(self, section: ReportSection) -> None:
        doc = self._doc
        doc.add_heading(section.title, level=1)
        if section.subtitle:
            run = doc.add_paragraph().add_run(section.subtitle)
            run.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        if section.meta:
            run = doc.add_paragraph().add_run("  ·  ".join(section.meta))
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        for block in section.blocks:
            kind = block.get("type")
            handler = getattr(self, f"_block_{kind}", None)
            if handler:
                handler(section, block)

    def _block_heading(self, section, block) -> None:
        level = block.get("level") or 2
        self._doc.add_heading(str(block.get("text", "")), level=min(max(level, 2), 4))

    def _block_pagebreak(self, section, block) -> None:
        self._doc.add_page_break()

    def _runs(self, paragraph, block: dict[str, Any], text: str, size: float | None = None) -> None:
        for style, content in paragraph_runs(block, text):
            run = paragraph.add_run(content)
            run.bold = style == "b" or None
            run.italic = style == "i" or None
            if style == "c":
                run.font.name = _MONO
                run.font.size = Pt(9.5)
            elif size:
                run.font.size = Pt(size)

    def _block_text(self, section, block) -> None:
        for text in block.get("paragraphs", []):
            self._runs(self._doc.add_paragraph(), block, str(text))

    def _block_list(self, section, block) -> None:
        style = "List Number" if block.get("ordered") else "List Bullet"
        for item in block.get("items", []):
            self._runs(self._doc.add_paragraph(style=style), {"markdown": True}, str(item))

    def _block_alert(self, section, block) -> None:
        fill, bar = _ALERT.get(block.get("kind"), _ALERT["info"])
        table = self._doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        _shade(cell, fill)
        _borders(cell, left=(bar, 24))
        cell.paragraphs[0].add_run(str(block.get("message", "")))
        self._doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _block_tiles(self, section, block) -> None:
        items = block.get("items", [])
        if not items:
            return
        columns = min(len(items), 4)
        rows = (len(items) + columns - 1) // columns
        table = self._doc.add_table(rows=rows, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for index, item in enumerate(items):
            cell = table.rows[index // columns].cells[index % columns]
            _shade(cell, tint(self._primary, 0.9))
            _borders(cell, top=("FFFFFF", 12), bottom=("FFFFFF", 12), left=("FFFFFF", 12), right=("FFFFFF", 12))
            label = cell.paragraphs[0]
            label.paragraph_format.space_after = Pt(0)
            run = label.add_run(str(item.get("label", "")))
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            value = cell.add_paragraph()
            value.paragraph_format.space_after = Pt(0)
            big = value.add_run(str(item.get("value", "")))
            big.bold = True
            big.font.size = Pt(13)
            if item.get("hint"):
                hint = cell.add_paragraph()
                hint.paragraph_format.space_after = Pt(0)
                small = hint.add_run(str(item["hint"]))
                small.font.size = Pt(7.5)
                small.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        self._doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _block_table(self, section, block) -> None:
        columns = [str(c) for c in block.get("columns", [])]
        rows = block.get("rows", [])
        if block.get("title"):
            title = self._doc.add_paragraph()
            title.paragraph_format.keep_with_next = True
            title.paragraph_format.space_before = Pt(8)
            run = title.add_run(str(block["title"]))
            run.bold = True
            run.font.size = Pt(10.5)
        if not columns:
            return
        size = 9 if len(columns) <= 5 else 8 if len(columns) <= 8 else 7
        table = self._doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        lengths = [
            max([len(columns[i])] + [len(cell_text(r[i])) for r in rows[:60] if i < len(r)])
            for i in range(len(columns))
        ]
        total = sum(min(n, 30) + 4 for n in lengths)
        widths = [Emu(int(self._width * (min(n, 30) + 4) / total)) for n in lengths]
        header = table.rows[0]
        _repeat_header(header)
        for index, name in enumerate(columns):
            cell = header.cells[index]
            cell.width = widths[index]
            _shade(cell, self._primary)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(name)
            run.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for row in rows:
            cells = table.add_row()
            _no_split(cells)
            for index in range(len(columns)):
                value = row[index] if index < len(row) else None
                text = cell_text(value)
                cell = cells.cells[index]
                cell.width = widths[index]
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                if index > 0 and is_numeric_text(text):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = paragraph.add_run(text)
                run.font.size = Pt(size)
                tone = cell_tone(value)
                if tone:
                    run.bold = tone == "strong" or None
                    run.font.color.rgb = RGBColor.from_string(_TONE_COLOR[tone])
        if block.get("note"):
            note = self._doc.add_paragraph().add_run(str(block["note"]))
            note.italic = True
            note.font.size = Pt(8.5)
            note.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        else:
            self._doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _block_figure(self, section, block) -> None:
        png = figure_image(section, block)
        paragraph = self._doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = bool(block.get("caption"))
        if not png or not (size := image_size(png)):
            paragraph.add_run("(Gráfico no disponible)").italic = True
            return
        width = min(self._width, Emu(int(Mm(170))))
        paragraph.add_run().add_picture(io.BytesIO(flatten_png(png)), width=width)
        _ = size
        if block.get("caption"):
            caption = self._doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(str(block["caption"]))
            run.italic = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    def _mono_block(self, source: str, fill: str, label: str | None = None) -> None:
        table = self._doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        _shade(cell, fill)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        lines = (source or "").replace("\t", "    ").split("\n")
        if label:
            tag = paragraph.add_run(label + "\n")
            tag.font.size = Pt(7.5)
            tag.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        for number, line in enumerate(lines):
            run = paragraph.add_run(line)
            run.font.name = _MONO
            run.font.size = Pt(8.5)
            if number < len(lines) - 1:
                run.add_break(WD_BREAK.LINE)
        self._doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _block_code(self, section, block) -> None:
        self._mono_block(str(block.get("source", "")), "F1F5F9", block.get("label"))

    def _block_output(self, section, block) -> None:
        self._mono_block(
            plain(str(block.get("text", ""))) if block.get("markdown") else str(block.get("text", "")), "FAFAF9"
        )
